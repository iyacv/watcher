import hashlib
import json
from datetime import datetime
from pathlib import Path
import xml.etree.ElementTree as ET


_LOG360_SEVERITIES = {"error", "failure", "warning", "information", "success"}


def _normalize_severity(raw: str) -> str:
    """Keep Log360's native vocabulary. F5 uses CVSS-derived severity;
    Log360 uses Windows event outcomes — they measure different things, so
    we don't try to flatten one into the other."""
    s = (raw or "").strip().lower()
    # Common synonyms that show up in Log360 exports
    aliases = {
        "audit success":   "success",
        "audit failure":   "failure",
        "info":            "information",
        "informational":   "information",
        "err":             "error",
        "warn":            "warning",
    }
    s = aliases.get(s, s)
    return s if s in _LOG360_SEVERITIES else "information"


def _make_hash(host: str, event_id: str, timestamp: str) -> str:
    key = f"{host}|{event_id}|{timestamp}"
    return hashlib.sha256(key.encode()).hexdigest()


def _text(el, *tags) -> str:
    for tag in tags:
        child = el.find(tag)
        if child is not None and child.text:
            return child.text.strip()
    return ""


def _record(event_id, timestamp, host, event_type, severity, user, display_name):
    sev = _normalize_severity(severity)
    ts  = timestamp or datetime.utcnow().isoformat()
    return {
        "source":       "log360",
        "host":         host.strip(),
        "event_id":     event_id.strip(),
        "event_type":   event_type.strip(),
        "severity":     sev,
        "user":         user.strip() if user else "",
        "display_name": display_name.strip() if display_name else "",
        "detected_at":  ts,
        "raw_hash":     _make_hash(host, event_id, ts),
    }


def _parse_xml(filepath: str) -> list:
    tree = ET.parse(filepath)
    root = tree.getroot()
    records = []

    events = root.findall(".//Event") + root.findall(".//event")
    for ev in events:
        records.append(_record(
            event_id     = _text(ev, "EventID", "eventId", "event_id"),
            timestamp    = _text(ev, "Timestamp", "timestamp", "Time", "time"),
            host         = _text(ev, "SourceHost", "sourceHost", "source_host", "Host", "host") or "unknown",
            event_type   = _text(ev, "EventType", "eventType", "event_type", "Type", "type"),
            severity     = _text(ev, "Severity", "severity") or "low",
            user         = _text(ev, "User", "user", "Username", "username"),
            display_name = _text(ev, "DisplayName", "displayName", "display_name", "Description", "description", "Message", "message"),
        ))
    return records


def _parse_json(filepath: str) -> list:
    with open(filepath, "r", encoding="utf-8") as f:
        data = json.load(f)

    events = data if isinstance(data, list) else data.get("events", data.get("Events", []))
    records = []

    for ev in events:
        records.append(_record(
            event_id     = str(ev.get("event_id") or ev.get("EventID") or ev.get("eventId") or ""),
            timestamp    = ev.get("timestamp") or ev.get("Timestamp") or ev.get("time") or "",
            host         = ev.get("source_host") or ev.get("SourceHost") or ev.get("host") or "unknown",
            event_type   = ev.get("event_type") or ev.get("EventType") or ev.get("type") or "",
            severity     = ev.get("severity") or ev.get("Severity") or "low",
            user         = ev.get("user") or ev.get("User") or ev.get("username") or "",
            display_name = ev.get("display_name") or ev.get("DisplayName") or ev.get("description") or ev.get("Description") or ev.get("message") or "",
        ))
    return records


EXPECTED_HEADERS = ("time", "log source", "event id", "display name", "source", "severity")


def _parse_xlsx(filepath: str) -> list:
    from openpyxl import load_workbook

    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    records = []

    rows = list(ws.iter_rows(values_only=True))

    i = 0
    while i < len(rows):
        cell_a = rows[i][0] if rows[i] else None
        if isinstance(cell_a, str) and cell_a.strip().lower() == "all events":
            # Find the next non-blank row — that should be the header
            j = i + 1
            while j < len(rows) and (not rows[j] or rows[j][0] is None or str(rows[j][0]).strip() == ""):
                j += 1
            if j >= len(rows):
                break

            header = rows[j]
            header_norm = tuple(
                (str(header[k]).strip().lower() if k < len(header) and header[k] is not None else "")
                for k in range(6)
            )
            if header_norm != EXPECTED_HEADERS:
                # Not the right "All Events" block — keep scanning
                i += 1
                continue

            # Read data rows until column A is empty
            k = j + 1
            while k < len(rows):
                row = rows[k]
                if not row or row[0] is None or str(row[0]).strip() == "":
                    break
                ts          = str(row[0]).strip() if row[0] is not None else ""
                log_source  = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
                event_id    = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
                display     = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
                source      = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ""
                severity    = str(row[5]).strip() if len(row) > 5 and row[5] is not None else "info"

                records.append(_record(
                    event_id     = event_id,
                    timestamp    = ts,
                    host         = log_source or display or "unknown",
                    event_type   = source,
                    severity     = severity,
                    user         = "",
                    display_name = display,
                ))
                k += 1
            break
        i += 1

    wb.close()
    return records


def parse_log360(filepath: str) -> list:
    ext = Path(filepath).suffix.lower()
    if ext == ".json":
        return _parse_json(filepath)
    if ext in (".xlsx", ".xlsm"):
        return _parse_xlsx(filepath)
    return _parse_xml(filepath)
